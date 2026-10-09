"""In-AMS written assignments: live writing, groups, integrity, and downloads."""
from datetime import datetime, timedelta
from io import BytesIO
from unittest.mock import patch

from docx import Document

from blueprints.class_management.models import ClassStudent, StudentNotification
from blueprints.student_management.models import Student as StudentMgmt
from blueprints.written_assignment.export import build_docx, build_pdf
from blueprints.written_assignment.models import (
    WrittenAssignment,
    WrittenAssignmentEvaluation,
    WrittenAssignmentPart,
    WrittenAssignmentPost,
    WrittenAssignmentPostFile,
    WrittenAssignmentPostVote,
    evaluate_paste,
    part_integrity,
)
from utils.ai.assignment_eval import scale_to_marks
from utils.outline_richtext import sanitize_outline_html
from extensions import db
from tests.conftest import TEST_PASSWORD, login_client
from user_models import User


AI_PARAGRAPH = (
    'Furthermore, it is important to note that the aforementioned framework plays a crucial role. '
    'Moreover, from a broader perspective, this multifaceted landscape underscores the pivotal role '
    'of comprehensive understanding. In conclusion, it can be argued that the intricate relationship '
    'remains a double-edged sword in today\'s world and therefore deserves a careful reading.'
)


def _ensure_student_user(username='S001', name='Student One'):
    user = User.query.filter_by(username=username).first()
    if not user:
        user = User(
            username=username,
            email=f'{username.lower()}@test.example',
            full_name=name,
            role='student',
        )
        user.set_password(TEST_PASSWORD)
        db.session.add(user)
        db.session.commit()
    return user


def _enroll(session_id, teacher_id, student_id, name):
    if not StudentMgmt.query.filter_by(student_id=student_id).first():
        db.session.add(StudentMgmt(student_id=student_id, name=name))
    if not ClassStudent.query.filter_by(session_id=session_id, student_id=student_id).first():
        db.session.add(ClassStudent(
            student_id=student_id,
            name=name,
            session_id=session_id,
            teacher_id=teacher_id,
        ))
    db.session.commit()


def _create(app, client, session_id, **overrides):
    data = {
        'title': 'Judicial review essay',
        'instructions': '<p>Use your own words.</p>',
        'mode': 'individual',
        'marks_out_of': '10',
        'paste_mode': 'disabled',
    }
    data.update(overrides)
    rv = client.post(f'/written-assignment/course/{session_id}', data=data, follow_redirects=True)
    assert rv.status_code == 200, rv.get_data(as_text=True)
    with app.app_context():
        assignment = WrittenAssignment.query.filter_by(session_id=session_id, title=data['title']).first()
        assert assignment is not None
        return assignment.id


def _save(client, assignment_id, html, **extra):
    payload = {'html': html, 'typed_chars': len(html), 'compose_ms': 4000}
    payload.update(extra)
    return client.post(
        f'/written-assignment/student/{assignment_id}/save',
        json=payload,
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )


def test_course_button_and_student_card(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    course = client.get(f'/written-assignment/course/{session_id}')
    assert course.status_code == 200
    assert 'New assignment' in course.get_data(as_text=True)

    from pathlib import Path
    index_html = Path('blueprints/class_management/templates/class_management/index.html').read_text()
    assert '/written-assignment/course/' in index_html
    assert '</i>Assignment' in index_html

    with app.app_context():
        _ensure_student_user()
    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    dash = client.get('/student/dashboard')
    assert dash.status_code == 200
    dash_html = dash.get_data(as_text=True)
    assert '>Assignments<' in dash_html
    assert '/written-assignment/student' in dash_html


def test_student_writes_and_teacher_sees_it_live(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user()
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id)

    with app.app_context():
        notes = StudentNotification.query.filter_by(type='assignment').all()
        assert notes
        assert any('Judicial review essay' in note.title for note in notes)

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    write = client.get(f'/written-assignment/student/{assignment_id}')
    assert write.status_code == 200
    write_html = write.get_data(as_text=True)
    assert 'id="assignmentEditor"' in write_html
    assert 'tinymce' in write_html
    assert 'blocks fontfamily fontsize' in write_html

    saved = _save(client, assignment_id, '<p>The <strong>court</strong> may review the decision.</p>')
    assert saved.status_code == 200, saved.get_data(as_text=True)
    body = saved.get_json()
    assert body['ok'] is True
    assert body['word_count'] >= 5

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    live = client.get(f'/written-assignment/{assignment_id}/live')
    assert live.status_code == 200
    payload = live.get_json()
    joined = ' '.join(row['html'] for row in payload['rows'])
    assert 'court' in joined
    assert '<strong>' in joined
    manage = client.get(f'/written-assignment/{assignment_id}')
    assert 'Student progress' in manage.get_data(as_text=True)
    assert 'court' in manage.get_data(as_text=True)


def test_paste_off_focus_and_ai_style(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user()
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, paste_mode='disabled')

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    _save(client, assignment_id, '<p>A short typed opening.</p>')
    blocked = _save(
        client,
        assignment_id,
        '<p>' + ('pasted sentence ' * 30) + '</p>',
        typed_chars=0,
        pasted_chars=400,
        paste_events=1,
        paste_words=60,
    )
    assert blocked.status_code == 400
    assert blocked.get_json()['code'] == 'paste_blocked'

    flagged = _save(client, assignment_id, f'<p>{AI_PARAGRAPH}</p>', typed_chars=400, compose_ms=20000)
    assert flagged.status_code == 200
    assert flagged.get_json()['style_flag'] is True
    assert flagged.get_json()['warning']

    event = client.post(
        f'/written-assignment/student/{assignment_id}/integrity',
        json={'event': 'focus_loss'},
        headers={'X-Requested-With': 'XMLHttpRequest'},
    )
    assert event.status_code == 200
    shot = client.post(
        f'/written-assignment/student/{assignment_id}/integrity',
        json={'event': 'screenshot_hint'},
        headers={'X-Requested-With': 'XMLHttpRequest'},
    )
    assert shot.status_code == 200

    with app.app_context():
        part = WrittenAssignmentPart.query.filter_by(assignment_id=assignment_id, student_id='S001').one()
        info = part_integrity(part)
        assert info['style_flag'] is True
        assert part.focus_loss_count == 1
        assert part.screenshot_hint_count == 1
        assert 'pasted sentence' not in (part.html_body or '')


def test_draft_stays_editable_and_unsubmit_before_deadline(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user()
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, title='Draft and unsubmit')
    with app.app_context():
        assignment = WrittenAssignment.query.get(assignment_id)
        assignment.deadline = datetime.utcnow() + timedelta(days=3)
        db.session.commit()

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    page = client.get(f'/written-assignment/student/{assignment_id}')
    html = page.get_data(as_text=True)
    assert page.status_code == 200
    assert 'Save draft' in html
    assert 'min-height: 720px' in html
    assert _save(client, assignment_id, '<p>Day one notes.</p>').status_code == 200
    assert _save(client, assignment_id, '<p>Day two, still a draft.</p>').status_code == 200

    submitted = client.post(
        f'/written-assignment/student/{assignment_id}/submit',
        json={},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )
    assert submitted.status_code == 200
    assert submitted.get_json()['status'] == 'submitted'
    locked = _save(client, assignment_id, '<p>Should not replace a submission.</p>')
    assert locked.status_code == 403

    opened = client.post(
        f'/written-assignment/student/{assignment_id}/unsubmit',
        json={},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )
    assert opened.status_code == 200
    assert opened.get_json()['status'] == 'draft'
    again = _save(client, assignment_id, '<p>Edited after unsubmit.</p>')
    assert again.status_code == 200

    with app.app_context():
        assignment = WrittenAssignment.query.get(assignment_id)
        assignment.deadline = datetime.utcnow() - timedelta(days=1)
        part = WrittenAssignmentPart.query.filter_by(assignment_id=assignment_id, student_id='S001').one()
        part.status = WrittenAssignmentPart.STATUS_SUBMITTED
        db.session.commit()
    late = client.post(
        f'/written-assignment/student/{assignment_id}/unsubmit',
        json={},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )
    assert late.status_code == 403


def test_teacher_sees_draft_and_evaluates_after_submit(app, client, teacher_user, class_session_w1, windows):
    assert scale_to_marks(7.5, 20) == 15
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user()
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, title='Progress then evaluate')
    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    assert _save(client, assignment_id, '<p>Draft progress for the teacher.</p>').status_code == 200
    page = client.get(f'/written-assignment/student/{assignment_id}')
    assert 'Discussion' in page.get_data(as_text=True)

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    manage = client.get(f'/written-assignment/{assignment_id}')
    html = manage.get_data(as_text=True)
    assert manage.status_code == 200
    assert 'Draft progress for the teacher' in html
    assert 'Progress only' in html
    assert 'class="wa-evaluate" hidden' in html
    with app.app_context():
        part = WrittenAssignmentPart.query.filter_by(assignment_id=assignment_id, student_id='S001').one()
        part_id = part.id
    blocked = client.post(
        f'/written-assignment/parts/{part_id}/mark',
        data={'teacher_mark': '8'},
        follow_redirects=True,
    )
    assert 'Mark this student after they submit' in blocked.get_data(as_text=True)

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    posted = client.post(
        f'/written-assignment/{assignment_id}/discussion',
        json={'body': 'Is the introduction enough?'},
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    assert posted.status_code == 200
    assert posted.get_json()['messages'][0]['body'] == 'Is the introduction enough?'
    submitted = client.post(
        f'/written-assignment/student/{assignment_id}/submit',
        json={},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )
    assert submitted.status_code == 200

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    reply = client.post(
        f'/written-assignment/{assignment_id}/discussion',
        json={'body': 'Add one case.'},
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    bodies = [row['body'] for row in reply.get_json()['messages']]
    assert 'Is the introduction enough?' in bodies
    assert 'Add one case.' in bodies
    ready = client.get(f'/written-assignment/{assignment_id}')
    assert 'Evaluate with AI' in ready.get_data(as_text=True)
    assert 'class="wa-evaluate" hidden' not in ready.get_data(as_text=True)

    evaluation = {
        'scores': {'argument': 7.5, 'evidence': 6, 'structure': 7, 'clarity': 8, 'relevance': 7},
        'strengths': ['Clear opening'],
        'weaknesses': ['Thin evidence'],
        'summary': 'A workable answer.',
        'suggested_score': 7,
        'suggested_mark': 7,
        'marks_out_of': 10,
    }
    with patch('utils.ai.assignment_eval.evaluate_writing', return_value=evaluation):
        generated = client.post(
            f'/written-assignment/parts/{part_id}/evaluate',
            json={},
            headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
        )
    assert generated.status_code == 200
    assert generated.get_json()['success'] is True
    with app.app_context():
        stored = WrittenAssignmentEvaluation.query.filter_by(target_type='part', target_id=part_id).one()
        assert 'A workable answer' in stored.body_json

    outsider = _ensure_user_outside(app)
    login_client(client, outsider, role='student', window_id=windows['w1_id'])
    denied = client.post(
        f'/written-assignment/{assignment_id}/discussion',
        json={'body': 'I am not in this course.'},
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    assert denied.status_code == 403


def test_deadline_blocks_until_teacher_reopens(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user()
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id)
    with app.app_context():
        assignment = WrittenAssignment.query.get(assignment_id)
        assignment.deadline = datetime.utcnow() - timedelta(days=2)
        db.session.commit()

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    closed = _save(client, assignment_id, '<p>Too late.</p>')
    assert closed.status_code == 403

    with app.app_context():
        part = WrittenAssignmentPart.query.filter_by(assignment_id=assignment_id, student_id='S001').one()
        part_id = part.id
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    reopened = client.post(f'/written-assignment/parts/{part_id}/reopen', follow_redirects=True)
    assert reopened.status_code == 200

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    again = _save(client, assignment_id, '<p>Late work allowed by the teacher.</p>')
    assert again.status_code == 200


def test_group_parts_are_separate_and_download_together(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user('S001', 'Student One')
        _ensure_student_user('S002', 'Student Two')
        _enroll(session_id, teacher_user['teacher_id'], 'S002', 'Student Two')

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, title='Group memo', mode='group', paste_mode='unlimited')
    created = client.post(
        f'/written-assignment/{assignment_id}/groups',
        data={'group_name': 'Bench A', 'student_ids': ['S001', 'S002']},
        follow_redirects=True,
    )
    assert created.status_code == 200
    assert 'Bench A' in created.get_data(as_text=True)

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    page = client.get(f'/written-assignment/student/{assignment_id}')
    assert 'Your part' in page.get_data(as_text=True)
    assert _save(client, assignment_id, '<p>S001 wrote the facts.</p>').status_code == 200

    login_client(client, 'S002', role='student', window_id=windows['w1_id'])
    assert _save(client, assignment_id, '<p>S002 wrote the argument.</p>').status_code == 200
    peers = client.get(f'/written-assignment/student/{assignment_id}/peers')
    rows = peers.get_json()['sections']
    peer_html = ' '.join(row['html'] for row in rows if not row['mine'])
    own_html = ' '.join(row['html'] for row in rows if row['mine'])
    assert 'S001 wrote the facts' in peer_html
    assert 'S002 wrote the argument' in own_html

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    docx = client.get(f'/written-assignment/student/{assignment_id}/download/docx')
    assert docx.status_code == 200
    assert 'wordprocessingml' in docx.mimetype
    document = Document(BytesIO(docx.data))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'S001' in text and 'S002' in text
    assert 'facts' in text and 'argument' in text

    outsider = _ensure_user_outside(app)
    login_client(client, outsider, role='student', window_id=windows['w1_id'])
    denied = client.post(
        f'/written-assignment/student/{assignment_id}/save',
        json={'html': '<p>No.</p>'},
        headers={'Accept': 'application/json'},
    )
    assert denied.status_code == 403


def test_group_members_get_separate_marks(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user('S001', 'Student One')
        _ensure_student_user('S002', 'Student Two')
        _enroll(session_id, teacher_user['teacher_id'], 'S002', 'Student Two')

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, title='Per-member marks', mode='group', paste_mode='unlimited')
    client.post(
        f'/written-assignment/{assignment_id}/groups',
        data={'group_name': 'Bench C', 'student_ids': ['S001', 'S002']},
        follow_redirects=True,
    )

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    assert client.get(f'/written-assignment/student/{assignment_id}').status_code == 200
    assert _save(client, assignment_id, '<p>S001 writing.</p>').status_code == 200
    submitted = client.post(
        f'/written-assignment/student/{assignment_id}/submit',
        json={},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )
    assert submitted.status_code == 200, submitted.get_json()

    login_client(client, 'S002', role='student', window_id=windows['w1_id'])
    assert client.get(f'/written-assignment/student/{assignment_id}').status_code == 200
    assert _save(client, assignment_id, '<p>S002 writing.</p>').status_code == 200
    submitted2 = client.post(
        f'/written-assignment/student/{assignment_id}/submit',
        json={},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )
    assert submitted2.status_code == 200, submitted2.get_json()

    with app.app_context():
        p1 = WrittenAssignmentPart.query.filter_by(assignment_id=assignment_id, student_id='S001').one()
        p2 = WrittenAssignmentPart.query.filter_by(assignment_id=assignment_id, student_id='S002').one()
        p1_id, p2_id = p1.id, p2.id

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    manage = client.get(f'/written-assignment/{assignment_id}')
    html = manage.get_data(as_text=True)
    assert 'Mark for Student One' in html
    assert 'Mark for Student Two' in html
    assert 'Each member still gets their own mark' in html

    client.post(
        f'/written-assignment/parts/{p1_id}/mark',
        data={'teacher_mark': '9', 'teacher_comment': 'Strong intro'},
        follow_redirects=True,
    )
    client.post(
        f'/written-assignment/parts/{p2_id}/mark',
        data={'teacher_mark': '6.5', 'teacher_comment': 'Needs detail'},
        follow_redirects=True,
    )
    with app.app_context():
        assert WrittenAssignmentPart.query.get(p1_id).teacher_mark == 9
        assert WrittenAssignmentPart.query.get(p2_id).teacher_mark == 6.5

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    page = client.get(f'/written-assignment/student/{assignment_id}').get_data(as_text=True)
    assert 'Your mark: 9' in page
    assert 'Strong intro' in page


def test_assignment_marks_fetch_into_empty_assessment(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user()
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, title='Fetchable essay', paste_mode='unlimited')

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    assert client.get(f'/written-assignment/student/{assignment_id}').status_code == 200
    assert _save(client, assignment_id, '<p>Marked essay.</p>').status_code == 200
    assert client.post(
        f'/written-assignment/student/{assignment_id}/submit',
        json={},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    ).status_code == 200
    with app.app_context():
        part = WrittenAssignmentPart.query.filter_by(assignment_id=assignment_id, student_id='S001').one()
        part.teacher_mark = 8
        db.session.commit()
        part_id = part.id

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    page = client.get(f'/class-management/assessment/{session_id}')
    html = page.get_data(as_text=True)
    assert page.status_code == 200
    assert 'Fetch marks from Assignment' in html
    assert 'Fetchable essay' in html

    fetched = client.post(
        f'/written-assignment/{assignment_id}/fetch-to-assessment',
        data={'target': 'assessment1'},
        follow_redirects=True,
    )
    assert fetched.status_code == 200
    with app.app_context():
        row = ClassStudent.query.filter_by(session_id=session_id, student_id='S001').one()
        assert row.assessment1 == 8

    blocked = client.post(
        f'/written-assignment/{assignment_id}/fetch-to-assessment',
        data={'target': 'assessment1'},
        follow_redirects=True,
    )
    assert 'already has marks' in blocked.get_data(as_text=True)
    with app.app_context():
        row = ClassStudent.query.filter_by(session_id=session_id, student_id='S001').one()
        row.assessment2 = 3
        db.session.commit()
    refused = client.post(
        f'/written-assignment/{assignment_id}/fetch-to-assessment',
        data={'target': 'assessment2'},
        follow_redirects=True,
    )
    assert 'already has marks' in refused.get_data(as_text=True)
    with app.app_context():
        row = ClassStudent.query.filter_by(session_id=session_id, student_id='S001').one()
        assert row.assessment2 == 3
        assert WrittenAssignmentPart.query.get(part_id).teacher_mark == 8


def test_group_sections_order_and_export(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user('S001', 'Student One')
        _ensure_student_user('S002', 'Student Two')
        _enroll(session_id, teacher_user['teacher_id'], 'S002', 'Student Two')

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, title='Sectioned memo', mode='group', paste_mode='unlimited')
    client.post(
        f'/written-assignment/{assignment_id}/groups',
        data={'group_name': 'Bench B', 'student_ids': ['S001', 'S002']},
        follow_redirects=True,
    )

    # S001 writes the introduction, S002 the middle, then S001 adds the conclusion.
    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    page = client.get(f'/written-assignment/student/{assignment_id}')
    html = page.get_data(as_text=True)
    assert 'Add section' in html and 'Group document' in html
    from blueprints.written_assignment.models import WrittenAssignmentSection
    with app.app_context():
        intro = WrittenAssignmentSection.query.join(
            WrittenAssignmentPart, WrittenAssignmentSection.part_id == WrittenAssignmentPart.id
        ).filter(
            WrittenAssignmentPart.assignment_id == assignment_id,
            WrittenAssignmentPart.student_id == 'S001',
        ).one()
        intro_id = intro.id
    saved = _save(client, assignment_id, '<p>The introduction.</p>', section_id=intro_id)
    assert saved.status_code == 200
    client.post(
        f'/written-assignment/student/{assignment_id}/sections/{intro_id}/rename',
        data={'title': 'Introduction'},
        follow_redirects=True,
    )

    login_client(client, 'S002', role='student', window_id=windows['w1_id'])
    client.get(f'/written-assignment/student/{assignment_id}')
    with app.app_context():
        middle = WrittenAssignmentSection.query.join(
            WrittenAssignmentPart, WrittenAssignmentSection.part_id == WrittenAssignmentPart.id
        ).filter(
            WrittenAssignmentPart.assignment_id == assignment_id,
            WrittenAssignmentPart.student_id == 'S002',
        ).one()
        middle_id = middle.id
    assert _save(client, assignment_id, '<p>The middle argument.</p>', section_id=middle_id).status_code == 200

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    added = client.post(
        f'/written-assignment/student/{assignment_id}/sections',
        data={'title': 'Conclusion'},
        follow_redirects=False,
    )
    assert added.status_code == 302
    with app.app_context():
        conclusion = WrittenAssignmentSection.query.filter_by(
            assignment_id=assignment_id, title='Conclusion'
        ).one()
        conclusion_id = conclusion.id
    assert _save(client, assignment_id, '<p>The conclusion.</p>', section_id=conclusion_id).status_code == 200

    peers = client.get(f'/written-assignment/student/{assignment_id}/peers').get_json()
    order = [row['title'] for row in peers['sections']]
    assert order == ['Introduction', 'Section 2', 'Conclusion']

    # Student moves the conclusion up, then back down.
    moved = client.post(
        f'/written-assignment/student/{assignment_id}/sections/move',
        json={'section_id': conclusion_id, 'direction': 'up'},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )
    assert moved.status_code == 200
    # The untitled middle section is auto-labelled by its current position.
    assert [row['title'] for row in moved.get_json()['sections']] == ['Introduction', 'Conclusion', 'Section 3']

    # Teacher moves it back down and sees the serial on the manage page.
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    back = client.post(
        f'/written-assignment/sections/{conclusion_id}/move',
        data={'direction': 'down'},
        follow_redirects=True,
    )
    assert back.status_code == 200
    manage_html = back.get_data(as_text=True)
    assert 'Section order' in manage_html

    docx = client.get(f'/written-assignment/{assignment_id}/download/docx?scope=all')
    assert docx.status_code == 200
    document = Document(BytesIO(docx.data))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert text.index('Introduction') < text.index('middle argument') < text.index('conclusion')
    assert '1. Introduction' in text and '3. Conclusion' in text


def _ensure_user_outside(app):
    with app.app_context():
        _ensure_student_user('S009', 'Outsider')
    return 'S009'


def test_docx_keeps_formatting_and_pdf_starts_when_engine_exists(app):
    sections = [{
        'heading': 'Student One (S001)',
        'subheading': 'draft · 4 words',
        'html': '<p>The <strong>court</strong> held that</p><ul><li>First point</li></ul>',
    }]
    raw = build_docx('Judicial review', ['LAW 4103'], sections)
    document = Document(BytesIO(raw))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'court' in text
    assert 'First point' in text
    assert 'Student One' in text

    with app.test_request_context():
        from flask import render_template
        from utils.integrity_highlight import PDF_HIGHLIGHT_CSS
        from utils.pdf_fonts import resolve_bangla_pdf_font, weasyprint_font_css
        css = weasyprint_font_css()
        assert 'data:font/ttf;base64,' in css
        assert 'PDFSerif' in css
        assert resolve_bangla_pdf_font() is not None
        assert 'Kalpurush' in css
        highlighted = [{
            'heading': 'Student One (S001)',
            'subheading': 'submitted · 12 words',
            'html': '<p class="hl-paste" style="background-color:#ffd54f">Pasted paragraph for review.</p>',
            'hl_classes': 'hl-paste-block',
            'mark_line': 'Teacher mark: 7.5 / 10',
        }]
        pdf_html = render_template(
            'written_assignment/export_pdf.html',
            title='Judicial review',
            meta_lines=['LAW 4103'],
            sections=highlighted,
            include_legend=True,
            highlight_css=PDF_HIGHLIGHT_CSS,
            kalpurush_uri=None,
        )
        assert 'background-color:#ffd54f' in pdf_html
        assert 'hl-paste-block' in pdf_html
        assert 'Teacher mark: 7.5 / 10' in pdf_html
        pdf = build_pdf('Judicial review', ['LAW 4103'], highlighted, include_legend=True)
    if pdf is not None:
        assert pdf[:4] == b'%PDF'
        assert b'CairoFont' not in pdf
        assert b'DroidSansFallback' not in pdf


def test_evaluate_paste_allows_typing_when_paste_is_off():
    assignment = WrittenAssignment(paste_disabled=True)
    part = WrittenAssignmentPart(paste_events_used=0, word_count=3)
    assert evaluate_paste(
        assignment, part,
        pasted_chars=0, paste_events=0, paste_words=0, old_words=3, new_words=40,
    ) is None
    assert evaluate_paste(
        assignment, part,
        pasted_chars=80, paste_events=1, paste_words=20, old_words=3, new_words=40,
    )


def test_discussion_like_mention_reply_and_attachment(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    png = (
        b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01'
        b'\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00'
        b'\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82'
    )
    with app.app_context():
        _ensure_student_user('S001', 'Student One')
        _ensure_student_user('S002', 'Student Two')
        _enroll(session_id, teacher_user['teacher_id'], 'S002', 'Student Two')

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, title='Discussion social')

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    posted = client.post(
        f'/written-assignment/{assignment_id}/discussion',
        json={'body': 'How should we open this paper?'},
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    assert posted.status_code == 200, posted.get_data(as_text=True)
    payload = posted.get_json()
    post_id = payload['messages'][0]['id']
    people_keys = {row['key'] for row in payload['people']}
    assert 'student:S001' in people_keys
    assert 'student:S002' in people_keys

    liked = client.post(
        f'/written-assignment/{assignment_id}/discussion/{post_id}/vote',
        json={'value': 1},
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    assert liked.status_code == 200
    first = liked.get_json()['messages'][0]
    assert first['likes'] == 1
    assert first['my_vote'] == 1

    liked_again = client.post(
        f'/written-assignment/{assignment_id}/discussion/{post_id}/vote',
        json={'value': 1},
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    assert liked_again.get_json()['messages'][0]['likes'] == 0

    disliked = client.post(
        f'/written-assignment/{assignment_id}/discussion/{post_id}/vote',
        json={'value': -1},
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    assert disliked.get_json()['messages'][0]['dislikes'] == 1

    login_client(client, 'S002', role='student', window_id=windows['w1_id'])
    reply = client.post(
        f'/written-assignment/{assignment_id}/discussion',
        data={
            'body': '@Student One here is a clipping.',
            'parent_id': str(post_id),
            'mentions': '["student:S001"]',
            'files': (BytesIO(png), 'clip.png'),
        },
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    assert reply.status_code == 200, reply.get_data(as_text=True)
    messages = reply.get_json()['messages']
    child = next(row for row in messages if row['parent_id'] == post_id)
    assert child['parent_name'] == 'Student One'
    assert any(item['key'] == 'student:S001' for item in child['mentions'])
    assert child['files'] and child['files'][0]['is_image'] is True
    file_url = child['files'][0]['url']
    image = client.get(file_url)
    assert image.status_code == 200
    assert image.data.startswith(b'\x89PNG')

    with app.app_context():
        note = StudentNotification.query.filter_by(type='assignment').order_by(StudentNotification.id.desc()).first()
        assert note is not None
        assert 'mentioned you' in note.title or 'replied to you' in note.title
        votes = WrittenAssignmentPostVote.query.all()
        assert len(votes) == 1
        assert votes[0].value == -1
        assert WrittenAssignmentPostFile.query.count() == 1
        assert WrittenAssignmentPost.query.count() == 2

    outsider = _ensure_user_outside(app)
    login_client(client, outsider, role='student', window_id=windows['w1_id'])
    denied = client.post(
        f'/written-assignment/{assignment_id}/discussion/{post_id}/vote',
        json={'value': 1},
        headers={'Accept': 'application/json', 'X-Requested-With': 'XMLHttpRequest'},
    )
    assert denied.status_code == 403
    blocked_file = client.get(file_url)
    assert blocked_file.status_code in (403, 302)


def test_highlighter_wraps_paste_and_ai_spans():
    from utils.integrity_highlight import highlight_integrity_html

    snippet = 'unique pasted clause about habeas corpus stays yellow'
    html = highlight_integrity_html(
        f'<p>Typed start. {snippet}. Typed end.</p><p>{AI_PARAGRAPH}</p>',
        paste_snippets=[snippet],
        ai_snippets=[AI_PARAGRAPH],
    )
    assert 'hl-paste' in html
    assert 'background-color:#ffd54f' in html
    assert snippet in html
    assert 'hl-ai' in html
    assert 'Furthermore' in html


def test_paste_ai_highlights_and_marked_pdf(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    snippet = 'This unique pasted clause about habeas corpus should glow yellow.'
    with app.app_context():
        _ensure_student_user()
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(
        app, client, session_id,
        title='Integrity highlight essay',
        paste_mode='unlimited',
    )

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    saved = _save(
        client,
        assignment_id,
        f'<p>A short typed intro. {snippet}</p><p>{AI_PARAGRAPH}</p>',
        typed_chars=40,
        pasted_chars=len(snippet),
        paste_events=1,
        paste_words=10,
        paste_snippets=[snippet],
        compose_ms=20000,
    )
    assert saved.status_code == 200, saved.get_data(as_text=True)
    submitted = client.post(
        f'/written-assignment/student/{assignment_id}/submit',
        json={},
        headers={'X-Requested-With': 'XMLHttpRequest', 'Accept': 'application/json'},
    )
    assert submitted.status_code == 200

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    live = client.get(f'/written-assignment/{assignment_id}/live')
    joined = ' '.join(row['html'] for row in live.get_json()['rows'])
    assert 'hl-paste' in joined
    assert snippet in joined
    assert 'hl-ai' in joined
    manage = client.get(f'/written-assignment/{assignment_id}')
    manage_html = manage.get_data(as_text=True)
    assert 'hl-paste' in manage_html
    assert 'PDF with marks' in manage_html

    with app.app_context():
        part = WrittenAssignmentPart.query.filter_by(assignment_id=assignment_id, student_id='S001').one()
        part_id = part.id
    marked = client.post(
        f'/written-assignment/parts/{part_id}/mark',
        data={'teacher_mark': '7.5', 'teacher_comment': 'Own words in the intro.'},
        follow_redirects=True,
    )
    assert marked.status_code == 200
    after = client.get(f'/written-assignment/{assignment_id}')
    after_html = after.get_data(as_text=True)
    assert '7.5' in after_html
    assert 'Own words in the intro.' in after_html

    pdf = client.get(
        f'/written-assignment/{assignment_id}/download/pdf?scope=student&student_id=S001&marked=1'
    )
    if pdf.status_code == 200:
        assert pdf.data[:4] == b'%PDF'
    else:
        assert pdf.status_code == 302


def test_sanitize_keeps_in_page_footnote_links():
    raw = (
        '<p>Seed law<sup class="wa-fn-ref"><a href="#wa-fn-1" id="wa-fnref-1">1</a></sup>.</p>'
        '<hr class="wa-fn-rule">'
        '<ol class="wa-footnotes"><li id="wa-fn-1">Rahman, A. (2020). <em>Seed sovereignty</em>.</li></ol>'
    )
    clean = sanitize_outline_html(raw)
    assert 'href="#wa-fn-1"' in clean
    assert 'id="wa-fn-1"' in clean
    assert 'id="wa-fnref-1"' in clean
    assert 'class="wa-fn-ref"' in clean
    assert 'class="wa-footnotes"' in clean
    assert 'target="_blank"' not in clean
    http = sanitize_outline_html('<p><a href="https://example.com">src</a></p>')
    assert 'target="_blank"' in http


def test_student_footnote_survives_save_and_docx(app, client, teacher_user, class_session_w1, windows):
    session_id = class_session_w1['session_id']
    with app.app_context():
        _ensure_student_user()
    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    assignment_id = _create(app, client, session_id, title='Footnote essay', paste_mode='unlimited')

    login_client(client, 'S001', role='student', window_id=windows['w1_id'])
    write = client.get(f'/written-assignment/student/{assignment_id}')
    html = write.get_data(as_text=True)
    assert 'wafootnote' in html
    assert 'Insert footnote' in html
    saved = _save(
        client,
        assignment_id,
        '<p>Farmers lost seed sovereignty'
        '<sup class="wa-fn-ref"><a href="#wa-fn-1" id="wa-fnref-1">1</a></sup>.</p>'
        '<hr class="wa-fn-rule">'
        '<ol class="wa-footnotes"><li id="wa-fn-1">Shiva, V. (2016). Who really feeds the world?</li></ol>',
    )
    assert saved.status_code == 200, saved.get_data(as_text=True)

    login_client(client, teacher_user['username'], window_id=windows['w1_id'])
    live = client.get(f'/written-assignment/{assignment_id}/live')
    joined = ' '.join(row['html'] for row in live.get_json()['rows'])
    assert 'wa-fn-ref' in joined
    assert 'Shiva' in joined
    assert 'href="#wa-fn-1"' in joined

    docx = client.get(
        f'/written-assignment/{assignment_id}/download/docx?scope=student&student_id=S001'
    )
    assert docx.status_code == 200
    document = Document(BytesIO(docx.data))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'Notes' in text
    assert 'Shiva' in text
    assert '1' in text
