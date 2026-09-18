"""Course-wise student feedback report for IQAC self-assessment.

Quantitative ratings are computed locally. The model writes narrative only,
using the supplied figures and verbatim comments.
"""
from __future__ import annotations

import json
import re
import statistics
from collections import Counter
from datetime import datetime
from io import BytesIO

from utils.ai.outline_parser import extract_json_from_response


LIKERT_SCORE = {
    'strongly disagree': 1,
    'disagree': 2,
    'neutral': 3,
    'agree': 4,
    'strongly agree': 5,
}

LIKERT_ITEMS = [
    ('A1', 'course_structure', 'section_a', 'Course content is structured in a comprehensible manner.'),
    ('A2', 'course_goals', 'section_a', 'The goals of the course are clear.'),
    ('A3', 'course_content_guidance', 'section_a', 'The course contents are explained in an understandable fashion.'),
    ('A4', 'course_interest', 'section_a', 'The course fosters my interest in the discussed topics.'),
    ('B1', 'course_plan_discussed', 'section_b', 'Course plan (assessment criteria/ content) was discussed in advance.'),
    ('B2', 'guidelines_received', 'section_b', 'Received oral instruction and written guidelines for continuous assessment.'),
    ('B3', 'assessment_helpful', 'section_b', 'Course projects/assignments/tests were helpful to demonstrate an understanding of the course material.'),
    ('B4', 'feedback_timely', 'section_b', 'Received feedback and grades of continuous assessments from course teacher in due time.'),
]

_EMAIL_RE = re.compile(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b')
_ID_RE = re.compile(
    r'\b(?:roll|id|student(?:\s*id)?|registration)\s*(?:no\.?|number|#)?\s*[:#]?\s*[A-Za-z0-9/.-]{3,}\b',
    re.I,
)
_PHONE_RE = re.compile(r'\b(?:\+?880|0)1[3-9]\d{8}\b')


def clip(text, limit=14000):
    text = (text or '').strip()
    if len(text) <= limit:
        return text
    return text[:limit] + '\n…[truncated]'


def likert_score(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        score = int(value)
        return score if 1 <= score <= 5 else None
    text = str(value).strip().lower()
    if not text:
        return None
    if text in LIKERT_SCORE:
        return LIKERT_SCORE[text]
    try:
        score = int(float(text))
    except (TypeError, ValueError):
        return None
    return score if 1 <= score <= 5 else None


def compute_item_statistics(scores):
    """Return n, mean, sd, % favourable (4–5), and distribution flags."""
    values = [int(v) for v in scores if v is not None]
    n = len(values)
    if not n:
        return {
            'n': 0,
            'mean': None,
            'sd': None,
            'favourable_pct': None,
            'distribution': {},
            'high_sd': False,
            'divided_opinion': False,
        }
    counts = Counter(values)
    mean = statistics.mean(values)
    sd = statistics.pstdev(values) if n == 1 else statistics.stdev(values)
    favourable = sum(1 for v in values if v >= 4)
    low = sum(1 for v in values if v <= 2)
    high = sum(1 for v in values if v >= 4)
    return {
        'n': n,
        'mean': round(mean, 2),
        'sd': round(sd, 2),
        'favourable_pct': round(100.0 * favourable / n, 1),
        'distribution': {str(k): counts.get(k, 0) for k in range(1, 6)},
        'high_sd': n >= 5 and sd >= 1.15,
        'divided_opinion': n >= 8 and (low / n) >= 0.25 and (high / n) >= 0.25,
    }


def anonymize_comment(text):
    cleaned = (text or '').strip()
    if not cleaned:
        return ''
    cleaned = _EMAIL_RE.sub('[email removed]', cleaned)
    cleaned = _PHONE_RE.sub('[phone removed]', cleaned)
    cleaned = _ID_RE.sub('[identifier removed]', cleaned)
    return cleaned


def parse_response_payload(row):
    raw = getattr(row, 'payload', None)
    if isinstance(raw, dict):
        return raw
    if not raw:
        return {}
    try:
        parsed = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _count_map(values):
    counts = Counter((str(v).strip() for v in values if str(v or '').strip()))
    return [{'label': label, 'count': count} for label, count in counts.most_common()]


def _missing(value):
    text = str(value).strip() if value is not None else ''
    return text if text else 'not reported'


def build_feedback_dataset(session, feedback_link, responses, enrolled_count=None):
    """Assemble PART A fields, item-wise stats, and anonymised comments."""
    from utils.tenant import current_tenant, year_is_postgraduate
    from utils.timezone import format_bd

    tenant = current_tenant()
    programme = tenant.default_program or 'not reported'
    if year_is_postgraduate(getattr(session, 'year', None)):
        options = tenant.program_options or ()
        programme = options[-1] if options else programme

    year_label = getattr(session, 'year', None) or 'not reported'
    term_label = getattr(session, 'term', None) or 'not reported'
    if year_label != 'not reported' and term_label != 'not reported':
        programme_and_level = f'{programme}, {year_label} Year, {term_label} Term'
    else:
        programme_and_level = programme

    teacher = getattr(session, 'teacher', None)
    teacher_name = getattr(teacher, 'name', None) if teacher is not None else None
    scope = (getattr(session, 'course_scope', None) or 'full').strip().lower()
    if scope in ('part_a', 'part_b') and teacher_name:
        teacher_name = f'{teacher_name} ({scope.replace("_", " ").title()})'

    credit = 'not reported'
    try:
        from blueprints.class_management.models import CourseOutline
        outline = CourseOutline.query.filter_by(session_id=session.id).first()
        if outline and outline.credit_value:
            credit = str(outline.credit_value).strip() or 'not reported'
    except Exception:
        credit = 'not reported'

    submitted_at = []
    item_scores = {code: [] for code, _key, _section, _label in LIKERT_ITEMS}
    methods = []
    study_times = []
    attendance = []
    effort = []
    comments = []

    for row in responses:
        payload = parse_response_payload(row)
        if getattr(row, 'submitted_at', None):
            submitted_at.append(row.submitted_at)
        section_a = payload.get('section_a') if isinstance(payload.get('section_a'), dict) else {}
        section_b = payload.get('section_b') if isinstance(payload.get('section_b'), dict) else {}
        section_c = payload.get('section_c') if isinstance(payload.get('section_c'), dict) else {}
        section_d = payload.get('section_d') if isinstance(payload.get('section_d'), dict) else {}
        sections = {'section_a': section_a, 'section_b': section_b}

        for code, key, section_name, _label in LIKERT_ITEMS:
            score = likert_score(sections.get(section_name, {}).get(key))
            if score is not None:
                item_scores[code].append(score)

        methods.extend(section_b.get('teaching_methods') or [])
        if section_c.get('study_time'):
            study_times.append(section_c.get('study_time'))
        if section_c.get('attendance_percent'):
            attendance.append(section_c.get('attendance_percent'))
        effort.extend(section_c.get('effort_focus') or [])

        comment = {}
        for field in ('likes', 'challenges', 'suggestions'):
            value = anonymize_comment(section_d.get(field) or '')
            if value:
                comment[field] = value
        if comment:
            comments.append(comment)

    items = []
    for code, _key, _section, label in LIKERT_ITEMS:
        stats = compute_item_statistics(item_scores[code])
        items.append({
            'item': code,
            'text': label,
            **stats,
        })

    scored_means = [row for row in items if row.get('mean') is not None]
    overall_mean = round(statistics.mean(row['mean'] for row in scored_means), 2) if scored_means else None
    ranked = sorted(scored_means, key=lambda row: (row['mean'], row['item']))
    lowest = ranked[:3]
    highest = list(reversed(ranked[-3:])) if ranked else []

    valid_n = len(responses)
    enrolled = enrolled_count if isinstance(enrolled_count, int) and enrolled_count >= 0 else None
    response_rate = None
    response_rate_label = 'not reported'
    if enrolled and enrolled > 0:
        response_rate = round(100.0 * valid_n / enrolled, 1)
        response_rate_label = f'{response_rate}% ({valid_n} of {enrolled})'
    indicative = bool(valid_n < 10 or (response_rate is not None and response_rate < 50))

    period = 'not reported'
    if submitted_at:
        first = min(submitted_at)
        last = max(submitted_at)
        start = format_bd(first, '%Y-%m-%d')
        end = format_bd(last, '%Y-%m-%d')
        period = start if start == end else f'{start} to {end}'

    instrument = (
        f'{len(LIKERT_ITEMS)}-item questionnaire, 5-point Likert scale '
        '(1 = strongly disagree to 5 = strongly agree), plus teaching-method and '
        'engagement items and three open-ended questions'
    )

    profile = [
        {'field': 'Course code and title', 'value': f"{_missing(getattr(session, 'course_code', None))} — {_missing(getattr(session, 'course_name', None))}"},
        {'field': 'Programme and level', 'value': programme_and_level},
        {'field': 'Term and academic session', 'value': f"{_missing(term_label)} Term, {_missing(getattr(session, 'academic_session', None) or getattr(session, 'year', None))}"},
        {'field': 'Course teacher(s)', 'value': _missing(teacher_name)},
        {'field': 'Credit hours', 'value': credit},
        {'field': 'Students enrolled', 'value': str(enrolled) if enrolled is not None else 'not reported'},
        {'field': 'Valid responses received', 'value': str(valid_n)},
        {'field': 'Response rate', 'value': response_rate_label},
        {'field': 'Survey instrument and scale', 'value': instrument},
        {'field': 'Mode of collection', 'value': 'Anonymous online form administered through the Academic Management System'},
        {'field': 'Period of collection', 'value': period},
        {'field': 'Form title', 'value': _missing(getattr(feedback_link, 'title', None))},
    ]

    return {
        'generated_at': datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC'),
        'discipline': tenant.display_with_university,
        'indicative_only': indicative,
        'profile': profile,
        'quantitative': {
            'likert_items': items,
            'overall_mean': overall_mean,
            'highest_rated': [{'item': row['item'], 'text': row['text'], 'mean': row['mean']} for row in highest],
            'lowest_rated': [{'item': row['item'], 'text': row['text'], 'mean': row['mean']} for row in lowest],
            'teaching_methods': _count_map(methods),
            'study_time': _count_map(study_times),
            'attendance': _count_map(attendance),
            'effort_focus': _count_map(effort),
            'scale': '1 = strongly disagree … 5 = strongly agree; favourable = ratings of 4 and 5',
        },
        'comments': comments,
        'comment_count': len(comments),
        'response_count': valid_n,
        'enrolled_count': enrolled,
        'response_rate': response_rate,
    }


SYSTEM_PROMPT = """You are assisting a university discipline in preparing a course-wise Student Feedback Report for institutional quality assurance and self-assessment purposes. The report will be read by the Institutional Quality Assurance Cell (IQAC) and may form part of a self-assessment exercise or accreditation submission. Write accordingly: formal, evidence-based, restrained in its claims, and auditable.

QUANTITATIVE ANALYSIS
- Use only the supplied item-wise statistics. Do not recalculate, invent, extrapolate, or estimate figures.
- If indicative_only is true (response rate below 50% or fewer than 10 valid responses), state clearly at the outset of the methodology and carry that qualification into the conclusions.
- Where an item is flagged high_sd or divided_opinion, say so — divided opinion is a finding and should not be flattened into an average.
- Do not apply inferential statistics or significance testing.

QUALITATIVE ANALYSIS
- Code the open comments inductively into themes. Do not force them into predetermined categories, but expect themes to cluster around: course content and coverage; teaching method and delivery; assessment and feedback; workload and pacing; reading materials and access to sources; classroom and physical facilities; and relevance to professional or further academic work.
- For each theme, state how many comments it draws on. Distinguish a view held by many respondents from a view expressed once.
- Quote sparingly and only where the wording is illustrative. Quotations must be verbatim from the supplied comments, marked as such, and stripped of anything identifying a student.
- Report negative and positive themes with equal care.

TRIANGULATION
- Compare ratings and comments. Where comments explain, contradict, or qualify a rating, say so explicitly.
- Where the two datasets diverge without explanation, record the divergence rather than resolving it.

EVIDENTIARY DISCIPLINE
- Every statement must be traceable to a specific item number, a stated frequency, or a quoted comment.
- If a field is "not reported" or a figure cannot be calculated, write "not reported" and list it under Limitations.
- Do not infer a cause the data does not support.

ANONYMITY AND PROPRIETY
- No student may be identifiable. Do not reproduce abusive or personal comments; record their substance in neutral terms if they bear on teaching or the course, and note the number excluded.

REPORT SECTIONS TO WRITE
2. Methodology — instrument, scale, mode and period of collection, response rate, and how comments were analysed. Three to four sentences.
3. Quantitative narrative — a brief reading of the supplied table. Do not simply restate the table in prose. No bullet points.
5. Triangulated key findings — four to six numbered findings drawing on both datasets. No bullet points in the finding text. Where a finding relates to curriculum content and structure, teaching-learning and assessment, student support services, or physical facilities and learning resources, name that area.
6. Strengths identified.
7. Areas requiring improvement — specific and diagnostic, not general.
8. Action plan — every entry in section 7 must appear here. Actions must fall within the authority of the course teacher, the Discipline, or the Academic Committee; flag referrals where resources are outside Discipline control.
9. Limitations.

STYLE
- Formal academic register. British spelling and Commonwealth conventions.
- Third person throughout. Avoid "we found"; prefer "the survey indicates".
- Evaluative language must be proportionate. Avoid "overwhelmingly", "excellent", "poor", and similar unanchored intensifiers.
- Return JSON only, matching the schema in the user message.
- Keep the JSON compact so it is not truncated: methodology 3–4 sentences; quantitative narrative one short paragraph; at most eight themes with two or three sentences each and at most one short quotation; four to six findings of two or three sentences; strengths, improvements and limitations as short bullet phrases; action-plan cells one sentence each. Do not paste the comment annex back into the JSON.
"""


def _text(value):
    return str(value or '').strip()


def _string_list(value):
    if isinstance(value, list):
        return [_text(item) for item in value if _text(item)]
    text = _text(value)
    return [text] if text else []


def _theme_list(value):
    themes = []
    if not isinstance(value, list):
        return themes
    for item in value:
        if not isinstance(item, dict):
            continue
        title = _text(item.get('theme') or item.get('title'))
        summary = _text(item.get('summary') or item.get('text'))
        if not title and not summary:
            continue
        count = item.get('comment_count') or item.get('frequency') or 0
        try:
            count = int(count)
        except (TypeError, ValueError):
            count = 0
        themes.append({
            'theme': title or 'Theme',
            'comment_count': count,
            'summary': summary,
            'quotation': _text(item.get('quotation') or item.get('quote')),
        })
    return themes


def _finding_list(value):
    findings = []
    if not isinstance(value, list):
        return findings
    for item in value:
        if isinstance(item, dict):
            text = _text(item.get('text') or item.get('finding'))
            if text:
                findings.append({
                    'text': text,
                    'qa_area': _text(item.get('qa_area') or item.get('quality_assurance_area')),
                })
        else:
            text = _text(item)
            if text:
                findings.append({'text': text, 'qa_area': ''})
    return findings


def _action_rows(value):
    rows = []
    if not isinstance(value, list):
        return rows
    for item in value:
        if not isinstance(item, dict):
            continue
        issue = _text(item.get('issue') or item.get('area'))
        action = _text(item.get('proposed_action') or item.get('action'))
        if not issue and not action:
            continue
        rows.append({
            'issue': issue,
            'proposed_action': action,
            'responsibility': _text(item.get('responsibility')),
            'timeline': _text(item.get('timeline')),
            'verification': _text(item.get('verification') or item.get('how_verified')),
        })
    return rows


INCOMPLETE_REPORT_MESSAGE = (
    'এআই রিপোর্ট অসম্পূর্ণ ছিল। আবার Generate চাপুন।'
)


def _clip_comment_field(value, limit=280):
    text = (value or '').strip()
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + '…'


def _prompt_comments(comments, limit=40):
    lines = []
    for index, comment in enumerate(comments[:limit], start=1):
        parts = []
        for key in ('likes', 'challenges', 'suggestions'):
            value = _clip_comment_field(comment.get(key) or '')
            if value:
                parts.append(f'{key}: {value}')
        if parts:
            lines.append(f'{index}. ' + ' | '.join(parts))
    omitted = len(comments) - limit
    if omitted > 0:
        lines.append(f'[{omitted} further comments omitted from the prompt]')
    return lines


def _prompt_ratings(quantitative):
    quantitative = quantitative or {}
    items = []
    for item in quantitative.get('likert_items') or []:
        items.append({
            'item': item.get('item'),
            'text': item.get('text'),
            'n': item.get('n'),
            'mean': item.get('mean'),
            'sd': item.get('sd'),
            'favourable_pct': item.get('favourable_pct'),
            'high_sd': bool(item.get('high_sd')),
            'divided_opinion': bool(item.get('divided_opinion')),
        })
    return {
        'scale': quantitative.get('scale'),
        'overall_mean': quantitative.get('overall_mean'),
        'highest_rated': quantitative.get('highest_rated') or [],
        'lowest_rated': quantitative.get('lowest_rated') or [],
        'likert_items': items,
        'teaching_methods': (quantitative.get('teaching_methods') or [])[:8],
        'study_time': quantitative.get('study_time') or [],
        'attendance': quantitative.get('attendance') or [],
    }


def generate_feedback_report(dataset):
    from utils.ai.client import generate_outline_json_with_meta, get_active_provider_setting

    get_active_provider_setting()

    user_payload = {
        'task': 'iqac_course_feedback_report',
        'discipline': dataset.get('discipline'),
        'indicative_only': dataset.get('indicative_only'),
        'course_and_survey_profile': dataset.get('profile'),
        'ratings_data': _prompt_ratings(dataset.get('quantitative')),
        'verbatim_student_comments': _prompt_comments(dataset.get('comments') or []),
        'output_schema': {
            'methodology': 'string',
            'quantitative_narrative': 'string',
            'qualitative_findings': [
                {'theme': 'string', 'comment_count': 0, 'summary': 'string', 'quotation': 'string'},
            ],
            'triangulated_key_findings': [{'text': 'string', 'qa_area': 'string'}],
            'strengths': ['string'],
            'areas_requiring_improvement': ['string'],
            'action_plan': [{
                'issue': 'string',
                'proposed_action': 'string',
                'responsibility': 'string',
                'timeline': 'string',
                'verification': 'string',
            }],
            'limitations': ['string'],
            'excluded_comments_count': 0,
            'excluded_comments_reason': 'string',
        },
    }
    meta = generate_outline_json_with_meta(
        SYSTEM_PROMPT,
        clip(json.dumps(user_payload, ensure_ascii=False), limit=18000),
        max_tokens=8192,
    )
    from utils.ai.session_utils import reset_db_session
    reset_db_session()
    try:
        payload = extract_json_from_response(meta.get('text') or '')
    except json.JSONDecodeError as exc:
        raise ValueError(INCOMPLETE_REPORT_MESSAGE) from exc
    if not isinstance(payload, dict):
        raise ValueError(INCOMPLETE_REPORT_MESSAGE)

    narrative = {
        'methodology': _text(payload.get('methodology')),
        'quantitative_narrative': _text(payload.get('quantitative_narrative')),
        'qualitative_findings': _theme_list(payload.get('qualitative_findings')),
        'triangulated_key_findings': _finding_list(payload.get('triangulated_key_findings')),
        'strengths': _string_list(payload.get('strengths')),
        'areas_requiring_improvement': _string_list(payload.get('areas_requiring_improvement')),
        'action_plan': _action_rows(payload.get('action_plan')),
        'limitations': _string_list(payload.get('limitations')),
        'excluded_comments_count': int(payload.get('excluded_comments_count') or 0),
        'excluded_comments_reason': _text(payload.get('excluded_comments_reason')),
    }
    if not narrative['methodology'] and not narrative['triangulated_key_findings']:
        raise ValueError(INCOMPLETE_REPORT_MESSAGE)

    report = dict(dataset)
    report['narrative'] = narrative
    report['provider'] = meta.get('provider')
    report['model_name'] = meta.get('model_name')
    return report


def load_stored_report(feedback_link):
    raw = getattr(feedback_link, 'ai_report_json', None) or ''
    if not raw:
        return None
    try:
        parsed = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return None
    return parsed if isinstance(parsed, dict) else None


def _set_run_font(run, name='Calibri', size=11, bold=False):
    from docx.shared import Pt
    from docx.oxml.ns import qn

    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def _add_paragraph(document, text, *, size=11, bold=False, space_after=8, style=None):
    from docx.shared import Pt

    para = document.add_paragraph(text, style=style) if style else document.add_paragraph(text)
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(0)
    if para.runs:
        _set_run_font(para.runs[0], size=size, bold=bold)
    return para


def _add_table(document, headers, rows):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(header)
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'D9E2F3')
        cell._tc.get_or_add_tcPr().append(shading)
        if cell.paragraphs and cell.paragraphs[0].runs:
            _set_run_font(cell.paragraphs[0].runs[0], size=10, bold=True)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = '' if value is None else str(value)
            if cell.paragraphs and cell.paragraphs[0].runs:
                _set_run_font(cell.paragraphs[0].runs[0], size=10)
    document.add_paragraph('')
    return table


def build_report_docx(report):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    from docx.oxml.ns import qn

    document = Document()
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    title = document.add_paragraph('Course-wise Student Feedback Report')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if title.runs:
        _set_run_font(title.runs[0], size=16, bold=True)

    subtitle = document.add_paragraph('Institutional Quality Assurance / Self-Assessment')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle.runs:
        _set_run_font(subtitle.runs[0], size=12, bold=True)

    _add_paragraph(document, report.get('discipline') or '', size=11)
    generated = report.get('generated_at') or ''
    if generated:
        _add_paragraph(document, f'Generated: {generated}', size=10)

    if report.get('indicative_only'):
        _add_paragraph(
            document,
            'Note: Findings are indicative rather than representative because the response rate is below 50 per cent or fewer than 10 valid responses were received.',
            size=11,
            bold=True,
        )

    _add_paragraph(document, '1. Course and Survey Profile', size=14, bold=True, space_after=6)
    profile_rows = [
        [row.get('field') or '', row.get('value') or '']
        for row in (report.get('profile') or [])
    ]
    _add_table(document, ['Field', 'Detail'], profile_rows)

    narrative = report.get('narrative') or {}
    _add_paragraph(document, '2. Methodology', size=14, bold=True, space_after=6)
    _add_paragraph(document, narrative.get('methodology') or 'not reported')

    _add_paragraph(document, '3. Quantitative Findings', size=14, bold=True, space_after=6)
    _add_paragraph(document, 'Table 1. Item-wise ratings (1 = strongly disagree to 5 = strongly agree).', size=10)
    item_rows = []
    for item in (report.get('quantitative') or {}).get('likert_items') or []:
        mean = item.get('mean')
        sd = item.get('sd')
        fav = item.get('favourable_pct')
        item_rows.append([
            item.get('item') or '',
            item.get('text') or '',
            item.get('n') if item.get('n') is not None else '—',
            f'{mean:.2f}' if isinstance(mean, (int, float)) else '—',
            f'{sd:.2f}' if isinstance(sd, (int, float)) else '—',
            f'{fav:.1f}%' if isinstance(fav, (int, float)) else '—',
        ])
    _add_table(
        document,
        ['Item', 'Statement', 'n', 'Mean', 'SD', '% favourable (4–5)'],
        item_rows,
    )
    overall = (report.get('quantitative') or {}).get('overall_mean')
    if overall is not None:
        _add_paragraph(document, f'Overall mean across items: {overall:.2f}.')
    _add_paragraph(document, narrative.get('quantitative_narrative') or '')

    _add_paragraph(document, '4. Qualitative Findings', size=14, bold=True, space_after=6)
    themes = narrative.get('qualitative_findings') or []
    if not themes:
        _add_paragraph(document, 'No open comments were available for thematic analysis.')
    for theme in themes:
        heading = theme.get('theme') or 'Theme'
        count = theme.get('comment_count')
        label = f'{heading} (n = {count})' if count else heading
        _add_paragraph(document, label, bold=True, space_after=4)
        _add_paragraph(document, theme.get('summary') or '')
        quote = theme.get('quotation') or ''
        if quote:
            _add_paragraph(document, f'“{quote}”', size=10)

    _add_paragraph(document, '5. Triangulated Key Findings', size=14, bold=True, space_after=6)
    findings = narrative.get('triangulated_key_findings') or []
    if not findings:
        _add_paragraph(document, 'not reported')
    for idx, finding in enumerate(findings, start=1):
        area = finding.get('qa_area') or ''
        suffix = f' ({area})' if area else ''
        _add_paragraph(document, f'{idx}. {finding.get("text") or ""}{suffix}')

    _add_paragraph(document, '6. Strengths Identified', size=14, bold=True, space_after=6)
    strengths = narrative.get('strengths') or []
    if strengths:
        for item in strengths:
            document.add_paragraph(item, style='List Bullet')
    else:
        _add_paragraph(document, 'not reported')

    _add_paragraph(document, '7. Areas Requiring Improvement', size=14, bold=True, space_after=6)
    areas = narrative.get('areas_requiring_improvement') or []
    if areas:
        for item in areas:
            document.add_paragraph(item, style='List Bullet')
    else:
        _add_paragraph(document, 'not reported')

    _add_paragraph(document, '8. Action Plan', size=14, bold=True, space_after=6)
    action_rows = [
        [
            row.get('issue') or '',
            row.get('proposed_action') or '',
            row.get('responsibility') or '',
            row.get('timeline') or '',
            row.get('verification') or '',
        ]
        for row in (narrative.get('action_plan') or [])
    ]
    if action_rows:
        _add_table(
            document,
            ['Issue', 'Proposed action', 'Responsibility', 'Timeline', 'How improvement will be verified'],
            action_rows,
        )
    else:
        _add_paragraph(document, 'not reported')

    _add_paragraph(document, '9. Limitations', size=14, bold=True, space_after=6)
    limitations = list(narrative.get('limitations') or [])
    excluded_n = narrative.get('excluded_comments_count') or 0
    excluded_reason = narrative.get('excluded_comments_reason') or ''
    if excluded_n:
        limitations.append(
            f'{excluded_n} comment(s) excluded under the propriety rule'
            + (f': {excluded_reason}' if excluded_reason else '.')
        )
    if limitations:
        for item in limitations:
            document.add_paragraph(item, style='List Bullet')
    else:
        _add_paragraph(document, 'not reported')

    _add_paragraph(document, '10. Annexes', size=14, bold=True, space_after=6)
    _add_paragraph(document, 'Annex A. Survey instrument (Likert items)', bold=True, space_after=4)
    _add_table(
        document,
        ['Item', 'Statement'],
        [[code, label] for code, _key, _section, label in LIKERT_ITEMS],
    )
    _add_paragraph(document, 'Annex B. Anonymised open comments', bold=True, space_after=4)
    comments = report.get('comments') or []
    if not comments:
        _add_paragraph(document, 'No open comments were recorded.')
    else:
        for idx, comment in enumerate(comments, start=1):
            parts = []
            if comment.get('likes'):
                parts.append(f'Liked: {comment["likes"]}')
            if comment.get('challenges'):
                parts.append(f'Challenges: {comment["challenges"]}')
            if comment.get('suggestions'):
                parts.append(f'Suggestions: {comment["suggestions"]}')
            _add_paragraph(document, f'{idx}. ' + ' | '.join(parts), size=10, space_after=6)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
