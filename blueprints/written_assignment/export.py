"""DOCX and PDF export for a written assignment (one student, one group, or everyone)."""
from html.parser import HTMLParser
from io import BytesIO

from utils.outline_richtext import html_is_empty


def _docx():
    """Import python-docx only when a file is actually built."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    return {
        'Document': Document,
        'WD_ALIGN_PARAGRAPH': WD_ALIGN_PARAGRAPH,
        'WD_BREAK': WD_BREAK,
        'OxmlElement': OxmlElement,
        'qn': qn,
        'Inches': Inches,
        'Pt': Pt,
        'RGBColor': RGBColor,
    }


def _apply_font(run, size=12):
    docx = _docx()
    run.font.size = docx['Pt'](size)
    run.font.name = 'Kalpurush'
    run.font.color.rgb = docx['RGBColor'](0x11, 0x18, 0x27)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(docx['qn']('w:rFonts'))
    if rFonts is None:
        rFonts = docx['OxmlElement']('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(docx['qn'](attr), 'Kalpurush')


class _DocxHtml(HTMLParser):
    """Turn sanitized assignment HTML into paragraphs, lists, and tables."""

    def __init__(self, doc):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.bold = 0
        self.italic = 0
        self.underline = 0
        self.strike = 0
        self.list_stack = []
        self.paragraph = None
        self.in_li = 0
        self.align = None
        self.heading = None
        self.in_table = 0
        self.rows = None
        self.cell = None
        self.super = 0

    def handle_starttag(self, tag, attrs):
        tag = (tag or '').lower()
        attrs = {((n or '').lower()): (v or '') for n, v in attrs}
        if tag in ('strong', 'b'):
            self.bold += 1
        elif tag in ('em', 'i'):
            self.italic += 1
        elif tag == 'u':
            self.underline += 1
        elif tag in ('s', 'strike'):
            self.strike += 1
        elif tag == 'br':
            self._run('').add_break()
        elif tag == 'hr':
            self._close_paragraph()
            rule = self.doc.add_paragraph()
            pt = _docx()['Pt']
            rule.paragraph_format.space_before = pt(6)
            rule.paragraph_format.space_after = pt(6)
        elif tag == 'sup':
            self.super += 1
        elif tag in ('ul', 'ol'):
            self._close_paragraph()
            classes = attrs.get('class') or ''
            if 'wa-footnotes' in classes:
                head = self.doc.add_paragraph()
                run = head.add_run('Notes')
                run.bold = True
                run.italic = True
                _apply_font(run, 12)
            self.list_stack.append({'type': tag, 'index': 0})
        elif tag == 'li':
            self._close_paragraph()
            self.in_li += 1
            if self.list_stack:
                self.list_stack[-1]['index'] += 1
            self._start_paragraph(list_item=True)
        elif tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            self._close_paragraph()
            self.heading = int(tag[1])
            self.align = _align_of(attrs)
            self._start_paragraph()
        elif tag in ('p', 'div', 'blockquote'):
            if self.in_li and self.paragraph is not None and not self._paragraph_has_text():
                self.align = _align_of(attrs) or self.align
                return
            self._close_paragraph()
            self.align = _align_of(attrs)
            self._start_paragraph(quote=(tag == 'blockquote'))
        elif tag == 'table':
            self._close_paragraph()
            self.in_table += 1
            if self.in_table == 1:
                self.rows = []
        elif tag == 'tr' and self.in_table == 1:
            self.rows.append([])
        elif tag in ('td', 'th') and self.in_table == 1 and self.rows is not None:
            self.cell = []
            self.rows[-1].append({'header': tag == 'th', 'text': self.cell})
        elif tag == 'a' and not self.in_table:
            href = (attrs.get('href') or '').strip()
            if href:
                self._pending_href = href

    def handle_endtag(self, tag):
        tag = (tag or '').lower()
        if tag in ('strong', 'b') and self.bold:
            self.bold -= 1
        elif tag in ('em', 'i') and self.italic:
            self.italic -= 1
        elif tag == 'u' and self.underline:
            self.underline -= 1
        elif tag in ('s', 'strike') and self.strike:
            self.strike -= 1
        elif tag == 'sup' and self.super:
            self.super -= 1
        elif tag in ('ul', 'ol') and self.list_stack:
            self._close_paragraph()
            self.list_stack.pop()
        elif tag == 'li' and self.in_li:
            self._close_paragraph()
            self.in_li -= 1
        elif tag in ('p', 'div', 'blockquote', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            self._close_paragraph()
            self.heading = None
            self.align = None
        elif tag == 'table' and self.in_table:
            self.in_table -= 1
            if self.in_table == 0:
                self._flush_table()
        elif tag in ('td', 'th') and self.in_table == 1:
            self.cell = None

    def handle_data(self, data):
        if not data:
            return
        if self.in_table == 1 and self.cell is not None:
            self.cell.append(data)
            return
        if self.in_table:
            return
        if self.paragraph is None:
            self._start_paragraph()
        self._run(data)

    def close(self):
        super().close()
        self._close_paragraph()

    def _paragraph_has_text(self):
        if self.paragraph is None:
            return False
        return any((run.text or '').strip() for run in self.paragraph.runs)

    def _start_paragraph(self, list_item=False, quote=False):
        style = None
        prefix = ''
        if list_item and self.list_stack:
            top = self.list_stack[-1]
            if top['type'] == 'ol':
                style = 'List Number'
            else:
                style = 'List Bullet'
        try:
            self.paragraph = self.doc.add_paragraph(style=style)
        except KeyError:
            self.paragraph = self.doc.add_paragraph()
            if list_item and self.list_stack:
                top = self.list_stack[-1]
                prefix = f"{top['index']}. " if top['type'] == 'ol' else '• '
        inches = _docx()['Inches']
        if quote:
            self.paragraph.paragraph_format.left_indent = inches(0.3)
        if self.align:
            self.paragraph.alignment = self.align
        if self.list_stack and list_item:
            self.paragraph.paragraph_format.left_indent = inches(0.25 * len(self.list_stack))
        if prefix:
            self._run(prefix)
        self._list_prefix_done = True

    def _run(self, text):
        if self.paragraph is None:
            self._start_paragraph()
        run = self.paragraph.add_run(text)
        size = 12
        if self.heading == 1:
            size = 18
        elif self.heading == 2:
            size = 16
        elif self.heading == 3:
            size = 14
        run.bold = bool(self.bold or (self.heading and self.heading <= 3))
        run.italic = bool(self.italic)
        run.underline = bool(self.underline)
        run.font.strike = bool(self.strike)
        if self.super:
            run.font.superscript = True
            size = 9
        _apply_font(run, size)
        return run

    def _close_paragraph(self):
        self.paragraph = None

    def _flush_table(self):
        rows = self.rows or []
        self.rows = None
        width = max((len(row) for row in rows), default=0)
        if not width:
            return
        table = self.doc.add_table(rows=len(rows), cols=width)
        try:
            table.style = 'Table Grid'
        except KeyError:
            pass
        for r_index, row in enumerate(rows):
            for c_index in range(width):
                cell = table.cell(r_index, c_index)
                payload = row[c_index] if c_index < len(row) else None
                text = ''
                header = False
                if payload:
                    text = ''.join(payload['text'])
                    header = payload['header']
                cell.text = ''
                para = cell.paragraphs[0]
                run = para.add_run(text.strip())
                run.bold = header
                _apply_font(run, 11)
        self.doc.add_paragraph()


def _align_map():
    align = _docx()['WD_ALIGN_PARAGRAPH']
    return {
        'left': align.LEFT,
        'center': align.CENTER,
        'right': align.RIGHT,
        'justify': align.JUSTIFY,
    }


def _align_of(attrs):
    mapping = _align_map()
    style = (attrs.get('style') or '').lower().replace(' ', '')
    for name, value in mapping.items():
        if f'text-align:{name}' in style:
            return value
    align = (attrs.get('align') or '').lower()
    return mapping.get(align)


def append_html(doc, html):
    text = (html or '').strip()
    if html_is_empty(text):
        para = doc.add_paragraph()
        run = para.add_run('No writing yet.')
        run.italic = True
        _apply_font(run, 12)
        return
    parser = _DocxHtml(doc)
    parser.feed(text)
    parser.close()


def build_docx(title, meta_lines, sections):
    """Return DOCX bytes. sections: [{heading, html}, ...]."""
    docx = _docx()
    doc = docx['Document']()
    normal = doc.styles['Normal']
    normal.font.name = 'Kalpurush'
    normal.font.size = docx['Pt'](12)
    heading = doc.add_heading(title or 'Assignment', level=0)
    for run in heading.runs:
        _apply_font(run, 20)
        run.bold = True
    for line in meta_lines or []:
        if not line:
            continue
        para = doc.add_paragraph()
        run = para.add_run(str(line))
        _apply_font(run, 11)
        run.font.color.rgb = docx['RGBColor'](0x37, 0x41, 0x51)
    for index, section in enumerate(sections or []):
        if index:
            breaker = doc.add_paragraph()
            breaker.add_run().add_break(docx['WD_BREAK'].PAGE)
        head = doc.add_heading(section.get('heading') or 'Part', level=1)
        for run in head.runs:
            _apply_font(run, 16)
            run.bold = True
        sub = section.get('subheading')
        if sub:
            para = doc.add_paragraph()
            run = para.add_run(sub)
            run.italic = True
            _apply_font(run, 11)
        append_html(doc, section.get('html') or '')
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def kalpurush_file_uri():
    try:
        from utils.pdf_fonts import resolve_bangla_pdf_font
        path = resolve_bangla_pdf_font()
    except Exception:
        return None
    return path.as_uri() if path else None


def build_pdf(title, meta_lines, sections, *, include_legend=False):
    """Return PDF bytes, or None when WeasyPrint cannot render."""
    from flask import render_template
    from utils.integrity_highlight import PDF_HIGHLIGHT_CSS
    from utils.pdf_fonts import inject_weasyprint_fonts

    html = render_template(
        'written_assignment/export_pdf.html',
        title=title or 'Assignment',
        meta_lines=[line for line in (meta_lines or []) if line],
        sections=sections or [],
        include_legend=include_legend,
        highlight_css=PDF_HIGHLIGHT_CSS,
        kalpurush_uri=kalpurush_file_uri(),
    )
    try:
        from weasyprint import HTML
    except Exception:
        return None
    try:
        html, base = inject_weasyprint_fonts(html)
        return HTML(string=html, base_url=base).write_pdf()
    except Exception:
        return None
